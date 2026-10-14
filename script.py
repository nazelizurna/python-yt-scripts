import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

FONT_NAME = "Times New Roman"
FONT_SIZE = 14  # points (Word uses pt, not px)
SEGMENTS_PER_PARAGRAPH = 6  # how many Whisper segments to merge into one paragraph


def download_youtube_mp3(url, output_filename="downloaded_audio"):
    import yt_dlp

    ydl_opts = {
        'format': 'bestaudio/best',
        'outtmpl': f'{output_filename}.%(ext)s',
        'noplaylist': True,
        # Use only a runtime you have installed (deno is the default;
        # remove this line if you have deno).
        'js_runtimes': {'node': {}},
        'remote_components': ['ejs:github'],
        'postprocessors': [{
            'key': 'FFmpegExtractAudio',
            'preferredcodec': 'mp3',
            'preferredquality': '192',
        }],
    }

    with yt_dlp.YoutubeDL(ydl_opts) as ydl:
        info = ydl.extract_info(url, download=True)
        title = info.get("title", "YouTube transcript")

    expected_file = f"{output_filename}.mp3"
    if not os.path.exists(expected_file):
        raise FileNotFoundError("Audio file extraction failed.")
    return expected_file, title


def transcribe(audio_path, model_name="base"):
    import whisper

    model = whisper.load_model(model_name)
    return model.transcribe(audio_path, fp16=False)  # fp16=False for CPU


def _set_run_font(run):
    run.font.name = FONT_NAME
    run.font.size = Pt(FONT_SIZE)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONT_NAME)


def _add_page_number_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    _set_run_font(run)
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def _add_paragraph(doc, text, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    _set_run_font(run)
    return p


def save_docx(result, title, url, output_path="youtube_transcript.docx"):
    doc = Document()

    # Default style: Times New Roman 14 pt
    style = doc.styles["Normal"]
    style.font.name = FONT_NAME
    style.font.size = Pt(FONT_SIZE)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)

    # Narrow margins (Word's "Narrow" preset = 1.27 cm / 0.5")
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Cm(1.27)
    section.left_margin = section.right_margin = Cm(1.27)

    # Page numbers, bottom centre
    _add_page_number_footer(section)

    # Title, then transcript without timestamps
    _add_paragraph(doc, title, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    segments = [s["text"].strip() for s in result["segments"] if s["text"].strip()]
    for i in range(0, len(segments), SEGMENTS_PER_PARAGRAPH):
        chunk = " ".join(segments[i:i + SEGMENTS_PER_PARAGRAPH])
        _add_paragraph(doc, chunk)

    doc.save(output_path)
    return output_path


if __name__ == "__main__":
    VIDEO_URL = "https://www.youtube.com/watch?v=doUeuzU-8SU"

    try:
        mp3_file, title = download_youtube_mp3(VIDEO_URL)
        result = transcribe(mp3_file, model_name="base")
        docx_path = save_docx(result, title, VIDEO_URL)
        print(f"\n--- Success ---\nSaved: {docx_path}")
        print(result["text"][:500])
    except Exception as e:
        print(f"Error: {e}")