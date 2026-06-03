from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


TEXT_ENCODINGS = ("utf-8-sig", "utf-8", "utf-16", "cp1251", "mac_cyrillic")


def read_text(path: Path) -> str:
    last_error: UnicodeDecodeError | None = None
    for encoding in TEXT_ENCODINGS:
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError as error:
            last_error = error
    if last_error is not None:
        raise last_error
    return path.read_text()


def set_run_font(run, *, bold: bool = False) -> None:
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()

    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")

    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"

    fld_char_sep = OxmlElement("w:fldChar")
    fld_char_sep.set(qn("w:fldCharType"), "separate")

    text = OxmlElement("w:t")
    text.text = "1"

    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")

    run._r.extend([fld_char_begin, instr_text, fld_char_sep, text, fld_char_end])
    set_run_font(run)


def add_text_paragraph(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Pt(18)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    set_run_font(run)


def convert_file(txt_path: Path, output_dir: Path) -> Path:
    document = Document()

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(14)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    section = document.sections[0]
    section.top_margin = Pt(72)
    section.bottom_margin = Pt(72)
    section.left_margin = Pt(72)
    section.right_margin = Pt(72)
    add_page_number(section.footer.paragraphs[0])

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    title.paragraph_format.space_after = Pt(12)
    title_run = title.add_run(txt_path.stem)
    set_run_font(title_run, bold=True)

    content = read_text(txt_path).replace("\r\n", "\n").replace("\r", "\n")
    for block in content.split("\n"):
        add_text_paragraph(document, block)

    output_path = output_dir / f"{txt_path.stem}.docx"
    document.save(output_path)
    return output_path


def main() -> None:
    parser = argparse.ArgumentParser(description="Convert .txt files in a folder to formatted .docx files.")
    parser.add_argument("directory", type=Path)
    parser.add_argument("--output-dir", type=Path)
    args = parser.parse_args()

    source_dir = args.directory.expanduser().resolve()
    output_dir = (args.output_dir or source_dir).expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    txt_files = sorted(source_dir.glob("*.txt"))
    for txt_path in txt_files:
        output_path = convert_file(txt_path, output_dir)
        print(output_path)

    print(f"Converted {len(txt_files)} file(s).")


if __name__ == "__main__":
    main()
